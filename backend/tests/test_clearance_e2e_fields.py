"""对照 WA318 样例关键数字。"""
import io

import openpyxl

from app.schemas.ledger import LedgerItemSchema, LedgerRecordResponse
from app.services.clearance_doc_service import ClearanceDocService


def wa318_record() -> LedgerRecordResponse:
    return LedgerRecordResponse(
        id=1,
        order_no="HT260720SZ",
        customer_code="WA318",
        consignee_name="CHEMZONE CO., LTD.",
        consignee_address="26/1D TRAN QUY CAP ST., BINH LOI TRUNG WARD, HO CHI MINH CITY, VIETNAM",
        destination="HOCHIMINH(CAT LAI),VIETNAM",
        loading_port="NANSHA,CHINA",
        price_term="CIF  HOCHIMINH",
        payment_terms="100% TT AFTER SHIPMENT 30 DAYS",
        pi_date="2026-07-20",
        currency="USD",
        items=[
            LedgerItemSchema(
                internal_code="HT-016H",
                product_en="FIXING AGENT HT-016H",
                quantity_kg=4000.0,
                unit_price=2.6,
                total_amount=10400.0,
                hs_code="340241",
                packaging_name="125kg /drum",
                drum_count=32,
                pallet_count=8,
                net_weight_kg=4000.0,
                gross_weight_kg=4308.0,
                volume_cbm=7.92,
            )
        ],
        status="saved",
    )


def _text(content: bytes) -> str:
    return "\n".join(
        str(c.value)
        for ws in openpyxl.load_workbook(io.BytesIO(content)).worksheets
        for row in ws.iter_rows()
        for c in row
        if c.value is not None
    )


def test_wa318_ci_pl_alignment():
    svc = ClearanceDocService()
    ov = {
        "container_no": "WHLU5694625",
        "seal_no": "WHA4051188",
        "invoice_date": "2026-08-01",
        "show_bank_block": True,
    }
    ci, _, _ = svc.generate("ci", wa318_record(), "honghao", ov)
    pl, _, _ = svc.generate("pl", wa318_record(), "honghao", ov)
    ci_text = _text(ci)
    pl_text = _text(pl)
    for t in (ci_text, pl_text):
        assert "WHLU5694625" in t
    # V2 §8.1：CI 无毛重/体积；N.W./G.W./CBM 在 PL（G.W./7.920 另见 SI）
    assert "4308" in pl_text or "4308.0" in pl_text
    assert "7.92" in pl_text or "7.920" in pl_text
    assert "IN260720SZ" in ci_text
    assert "PL260720SZ" in pl_text
    assert "TOTAL: 32 DRUMS PACKED ON 8 PALLETS" in ci_text
    assert "10400" in ci_text or "10400.0" in ci_text


def wa318_two_product_record() -> LedgerRecordResponse:
    """WA 风格双产品样例：同一 PI 下 CI/PL 合并出数。"""
    return LedgerRecordResponse(
        id=2,
        order_no="HT260720SZ",
        customer_code="WA318",
        consignee_name="CHEMZONE CO., LTD.",
        consignee_address="26/1D TRAN QUY CAP ST., BINH LOI TRUNG WARD, HO CHI MINH CITY, VIETNAM",
        destination="HOCHIMINH(CAT LAI),VIETNAM",
        loading_port="NANSHA,CHINA",
        price_term="CIF  HOCHIMINH",
        payment_terms="100% TT AFTER SHIPMENT 30 DAYS",
        pi_date="2026-07-20",
        currency="USD",
        items=[
            LedgerItemSchema(
                internal_code="HT-016H",
                product_en="FIXING AGENT HT-016H",
                quantity_kg=4000.0,
                unit_price=2.6,
                total_amount=10400.0,
                hs_code="340241",
                packaging_name="125kg /drum",
                drum_count=32,
                pallet_count=8,
                net_weight_kg=4000.0,
                gross_weight_kg=4308.0,
                volume_cbm=7.92,
            ),
            LedgerItemSchema(
                internal_code="HT-022S",
                product_en="SOFTENER HT-022S",
                quantity_kg=2000.0,
                unit_price=1.8,
                total_amount=3600.0,
                hs_code="340242",
                packaging_name="125kg /drum",
                drum_count=16,
                pallet_count=4,
                net_weight_kg=2000.0,
                gross_weight_kg=2154.0,
                volume_cbm=3.96,
            ),
        ],
        status="saved",
    )


def test_wa318_multi_item_ci_pl_alignment():
    """双产品 CI+PL 同票出数：descs / 合计 / 单号一致，无残留占位符。"""
    svc = ClearanceDocService()
    ov = {
        "container_no": "WHLU5694625",
        "seal_no": "WHA4051188",
        "invoice_date": "2026-08-01",
        "show_bank_block": True,
    }
    record = wa318_two_product_record()
    ci, _, _ = svc.generate("ci", record, "honghao", ov)
    pl, _, _ = svc.generate("pl", record, "honghao", ov)
    ci_text = _text(ci)
    pl_text = _text(pl)

    # 单号同源 PI HT260720SZ
    assert "IN260720SZ" in ci_text
    assert "PL260720SZ" in pl_text
    assert "HT260720SZ" in ci_text
    assert "HT260720SZ" in pl_text

    for t in (ci_text, pl_text):
        assert "FIXING AGENT HT-016H" in t
        assert "SOFTENER HT-022S" in t
        assert "{{" not in t

    # CI TOTAL = 4000+2000 qty / 10400+3600 amount
    ci_ws = openpyxl.load_workbook(io.BytesIO(ci)).worksheets[0]
    ci_total_qty = ci_total_amount = None
    for row in ci_ws.iter_rows(min_col=1, max_col=5):
        a, b, c, d, e = (cell.value for cell in row)
        if a == "TOTAL:" or b == "TOTAL:":
            ci_total_qty = float(c if c is not None else b)
            ci_total_amount = float(e)
    assert ci_total_qty == 6000.0
    assert ci_total_amount == 14000.0

    # PL TOTAL = net 6000 / gross 6462 / cbm 11.88 / packages(drums) 48
    pl_ws = openpyxl.load_workbook(io.BytesIO(pl)).worksheets[0]
    pl_total_pkg = pl_total_cbm = pl_total_net = pl_total_gross = None
    for row in pl_ws.iter_rows(min_col=1, max_col=6):
        a, b, c, d, e, f = (cell.value for cell in row)
        if a == "TOTAL:" or b == "TOTAL:":
            pl_total_pkg = float(c)
            pl_total_cbm = float(d)
            pl_total_net = float(e)
            pl_total_gross = float(f)
    assert pl_total_pkg == 48.0
    assert pl_total_cbm == round(7.92 + 3.96, 3)
    assert pl_total_net == 6000.0
    assert pl_total_gross == 6462.0
