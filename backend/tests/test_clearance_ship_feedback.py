# -*- coding: utf-8 -*-
"""船务反馈回归：HT20260808B01 三行（含无单价 LEVELLING AGENT）+ SI 补料表单。"""
import io

import openpyxl

from app.schemas.ledger import LedgerItemSchema, LedgerRecordResponse
from app.services.clearance_doc_service import ClearanceDocService


def chaity_record() -> LedgerRecordResponse:
    """截图同款：SILICONE SOFTENER + LEVELLING AGENT + LEVELLING AGENT(1开3 无单价)。"""
    return LedgerRecordResponse(
        id=735,
        order_no="HT20260808B01",
        customer_code="WA316",
        consignee_name="M/S CHAITY COMPOSITE LTD",
        consignee_address="CHOTTO SHILMONDI, TRIPURDI SONARGAON, NARAYAGONJ-1400, BANGLADESH",
        destination="Chattogram, Bangladesh",
        loading_port="Any Sea Port of CHINA",
        price_term="CNF CHATTOGRAM",
        payment_terms="100% LC 120 days from the date of BL",
        pi_date="2026-09-17",
        currency="USD",
        items=[
            LedgerItemSchema(
                internal_code="F-636-1",
                product_en="SILICONE SOFTENER",
                hs_code="3910000000",
                quantity_kg=5040.0,
                unit_price=2.4,
                total_amount=12096.0,
            ),
            LedgerItemSchema(
                internal_code="T28-EA",
                product_en="LEVELLING AGENT",
                hs_code="3402420000",
                quantity_kg=5040.0,
                unit_price=2.55,
                total_amount=12852.0,
            ),
            LedgerItemSchema(
                internal_code="T28-EA 1开3",
                product_en="LEVELLING AGENT",
                hs_code=None,
                quantity_kg=5040.0,
                unit_price=None,
                total_amount=None,
            ),
        ],
        status="pending",
    )


def _text(content: bytes) -> str:
    parts = []
    for ws in openpyxl.load_workbook(io.BytesIO(content)).worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    parts.append(str(cell.value))
    return "\n".join(parts)


def test_ci_missing_price_not_zero():
    """第三行无单价：价格/金额格留空，不打 0.0。"""
    svc = ClearanceDocService()
    content, _, _ = svc.generate("ci", chaity_record(), "honghao", {"show_bank_block": False})
    assert "{{" not in _text(content)

    ws = openpyxl.load_workbook(io.BytesIO(content)).worksheets[0]
    prices = {}
    amounts = {}
    for row in ws.iter_rows(min_col=1, max_col=5):
        a, b, c, d, e = (cell.value for cell in row)
        if isinstance(b, str) and (
            b.startswith("SILICONE SOFTENER")
            or b.startswith("LEVELLING AGENT")
        ) and not b.startswith("QTY") and not b.startswith("HS"):
            prices[b] = d
            amounts[b] = e

    third = [k for k in prices if "1开3" in k][0]
    assert prices[third] in (None, "")
    assert amounts[third] in (None, "")
    # 有价行正常
    assert any(str(prices[k]) in ("2.4", "2.40") for k in prices if "SILICONE" in k)


def test_ci_duplicate_product_en_distinguished():
    """同名 LEVELLING AGENT 用内部编号区分；唯一名保持干净。"""
    svc = ClearanceDocService()
    content, _, _ = svc.generate("ci", chaity_record(), "honghao", {})
    text = _text(content)
    assert "LEVELLING AGENT T28-EA" in text
    assert "LEVELLING AGENT T28-EA 1开3" in text
    # 唯一品名不硬加编号
    assert "SILICONE SOFTENER F-636-1" not in text
    assert "SILICONE SOFTENER" in text


def test_ci_multi_item_per_item_hs_block():
    """多品每行下面跟 QTY/HS 组，不是文末一条汇总 HS。"""
    svc = ClearanceDocService()
    content, _, _ = svc.generate("ci", chaity_record(), "honghao", {})
    text = _text(content)
    assert "HS CODE:3910000000" in text or "HS CODE: 3910000000" in text
    assert "HS CODE:3402420000" in text or "HS CODE: 3402420000" in text
    assert "QTY:5040" in text
    # 合计金额 = 12096+12852（第三行无价不计入）
    ws = openpyxl.load_workbook(io.BytesIO(content)).worksheets[0]
    total_amount = None
    total_qty = None
    for row in ws.iter_rows(min_col=1, max_col=5):
        a, b, c, d, e = (cell.value for cell in row)
        if a == "TOTAL:" or b == "TOTAL:":
            total_qty = float(c if c is not None else 0)
            if e is not None:
                total_amount = float(e)
    assert total_qty == 15120.0
    assert total_amount == 24948.0


def test_si_is_buliao_form_not_pl_shell():
    """SI 必须是补料表单：Shipper/Consignee/Notify/Marks/No of Packages。"""
    svc = ClearanceDocService()
    content, _, _ = svc.generate("si", chaity_record(), "honghao", {"vessel": "WAN HAI 289", "voyage": "S086"})
    text = _text(content)
    assert "补料" in text
    assert "Shipper" in text or "发货人" in text
    assert "Consignee" in text or "收货人" in text
    assert "Notify Party" in text or "通知人" in text
    assert "Marks" in text or "唛头" in text
    assert "No of Packages" in text or "数量" in text
    # 不再是 PL 壳的表头
    assert "PACKING QTY" not in text
    assert "BOOKING / SHIPPING INSTRUCTION" not in text


def test_si_desc_block_multi_product():
    """补料货描：多品编号列表 + 每品 HS/QTY。"""
    svc = ClearanceDocService()
    content, _, _ = svc.generate("si", chaity_record(), "honghao", {})
    text = _text(content)
    assert "1.SILICONE SOFTENER F-636-1" in text or "1.SILICONE SOFTENER" in text
    assert "2.LEVELLING AGENT" in text
    assert "3.LEVELLING AGENT" in text
    assert "H.S. Code" in text
    assert "QTY:5040KGS" in text
    assert "{{" not in text


def test_ci_no_stale_chemzone_contact():
    """公共模板不得带 WA318 样例写死的越南 TEL/EMAIL。"""
    svc = ClearanceDocService()
    content, _, _ = svc.generate("ci", chaity_record(), "honghao", {})
    text = _text(content)
    assert "THUY.PHAM@CHEMZONEVN.COM" not in text
    assert "+84 987 208 205" not in text
