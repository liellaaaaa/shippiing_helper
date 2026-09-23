import io

import openpyxl
import pytest

from app.schemas.ledger import LedgerItemSchema, LedgerRecordResponse
from app.services.clearance_doc_service import ClearanceDocService


def make_record() -> LedgerRecordResponse:
    return LedgerRecordResponse(
        id=1,
        order_no="HT260720SZ",
        customer_code="WA318",
        consignee_name="CHEMZONE CO., LTD.",
        consignee_address="26/1D TRAN QUY CAP ST., HO CHI MINH CITY, VIETNAM",
        consignee_tel="TAX CODE: 0316699088",
        destination="HOCHIMINH(CAT LAI), VIETNAM",
        loading_port="NANSHA, CHINA",
        price_term="CIF HOCHIMINH",
        payment_terms="100% TT AFTER SHIPMENT 30 DAYS",
        pi_date="2026-07-20",
        currency="USD",
        items=[
            LedgerItemSchema(
                internal_code="FIX-HT016H",
                product_en="FIXING AGENT HT-016H",
                quantity_kg=4000.0,
                unit_price=2.6,
                total_amount=10400.0,
                hs_code="340241",
                drum_count=32,
                pallet_count=8,
                net_weight_kg=4000.0,
                gross_weight_kg=4308.0,
                volume_cbm=7.92,
            )
        ],
        status="saved",
    )


@pytest.fixture
def svc():
    return ClearanceDocService()


def _all_text(wb) -> str:
    parts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    parts.append(str(cell.value))
    return "\n".join(parts)


def test_generate_ci(svc):
    content, doc_key, _ = svc.generate("ci", make_record(), "honghao", {
        "container_no": "WHLU5694625",
        "seal_no": "WHA4051188",
        "invoice_date": "2026-08-01",
        "show_bank_block": True,
    })
    assert doc_key.startswith("ci_")
    wb = openpyxl.load_workbook(io.BytesIO(content))
    text = _all_text(wb)
    assert "IN260720SZ" in text
    assert "HT260720SZ" in text
    assert "FIXING AGENT HT-016H" in text
    assert "WHLU5694625" in text
    assert "TOTAL 32 DRUMS PACKED ON 8 PALLETS ONLY." in text
    assert "{{" not in text
    assert "BENEFICIARY BANK" in text


def test_generate_ci_bank_off(svc):
    content, _, _ = svc.generate("ci", make_record(), "honghao", {"show_bank_block": False})
    text = _all_text(openpyxl.load_workbook(io.BytesIO(content)))
    assert "BENEFICIARY BANK" not in text


def test_generate_pl(svc):
    content, doc_key, _ = svc.generate("pl", make_record(), "honghao", {})
    assert doc_key.startswith("pl_")
    text = _all_text(openpyxl.load_workbook(io.BytesIO(content)))
    assert "PL260720SZ" in text
    assert "4308" in text or "4308.0" in text


def test_generate_coa_tbd_blank_pi(svc):
    content, doc_key, _ = svc.generate(
        "coa",
        make_record(),
        "honghao",
        {
            "batch_no": "SA20260626017",
            "expiry_rule": "plus_1y_minus_1d",
            "ph_label": "PH (1%)",
            "ph_spec": "6.5~7.5",
            "ph_result": "7.3",
            "solid_spec": "70±2",
            "solid_result": "70.72",
            "appearance_spec": "Yellow transparent mucus",
            "appearance_result": "Yellow transparent mucus",
        },
    )
    assert doc_key.startswith("coa_")
    text = _all_text(openpyxl.load_workbook(io.BytesIO(content)))
    assert "SA20260626017" in text
    assert "2026-06-26" in text
    assert "2027-06-25" in text
    assert "PH (1%)" in text
    assert "70.72" in text


def test_generate_si(svc):
    content, doc_key, _ = svc.generate(
        "si",
        make_record(),
        "honghao",
        {
            "vessel": "WAN HAI 289",
            "voyage": "S086",
            "bl_no": "GHOC26071803",
            "dest_agent_name": "STAR CONCORD (VIETNAM) CO., LTD.",
        },
    )
    assert doc_key.startswith("si_")
    text = _all_text(openpyxl.load_workbook(io.BytesIO(content)))
    assert "WAN HAI 289 / S086" in text
    assert "STAR CONCORD" in text
    assert "订舱补料" in text or "SHIPPING INSTRUCTION" in text


def test_unknown_doc_type(svc):
    with pytest.raises(ValueError):
        svc.generate("unknown", make_record(), "honghao", {})
