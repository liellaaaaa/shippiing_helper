# -*- coding: utf-8 -*-
"""清关四单生成：加载客户/公共模板 → 填占位符 → xlsx bytes。

对齐 MSDS generate_msds_from_template：
- 「标签：{{KEY}}」混排子串替换，保留样式
- 优先 customer_templates（一客一模板），否则公共模板
- overrides.extra_notes 追加说明行
"""
from __future__ import annotations

import base64
import copy
import re
import time
from io import BytesIO
from typing import Any, Dict, Optional, Tuple, Union

import openpyxl
from openpyxl.styles import Alignment

from app.schemas.ledger import LedgerRecordResponse
from app.services.clearance_fields import build_clearance_payload
from app.services import template_service

_DOC_TYPES = ("ci", "pl", "coa", "si")
_PLACEHOLDER_RE = re.compile(r"\{\{[^}]*\}\}")


def _fmt(value: Any) -> str:
    return "" if value is None else str(value)


def _fmt_num(value: Any) -> str:
    """数量/金额显示：整数去掉 .0，与成品票 QTY:5040KGS 一致。"""
    if value is None or value == "":
        return ""
    try:
        f = float(value)
    except (TypeError, ValueError):
        return str(value)
    if f == int(f) and abs(f) < 1e15:
        return str(int(f))
    return str(f)


def _coerce_number(text: str) -> Any:
    """填充后把纯数字串写回数值格，方便 Excel 求和/对齐成品票。"""
    s = (text or "").strip()
    if not s:
        return text
    try:
        if re.fullmatch(r"-?\d+", s):
            return int(s)
        if re.fullmatch(r"-?\d+\.\d+", s):
            return float(s)
    except ValueError:
        pass
    return text


def _extract_batch_test_fields(batch: Dict[str, Any]) -> Dict[str, str]:
    """从批次 tests[] /扁平字段提取 COA 检测项（报告写什么就填什么）。

    tests[] 形状：{key, name_cn, label_en, spec, result}
    扁平字段（appearance_spec / ph_label / …）优先于 tests[]（前端可编辑覆盖）。
    """
    out: Dict[str, str] = {
        "appearance_spec": "",
        "appearance_result": "",
        "ph_label": "",
        "ph_spec": "",
        "ph_result": "",
        "solid_label": "",
        "solid_spec": "",
        "solid_result": "",
        "odour_label": "",
        "odour_spec": "",
        "odour_result": "",
    }
    for t in batch.get("tests") or []:
        if not isinstance(t, dict):
            continue
        key = (t.get("key") or "").strip().lower()
        spec = t.get("spec") or ""
        result = t.get("result") or ""
        label = t.get("label_en") or ""
        if key == "appearance":
            # 外观规格全文保留（含「久置变深」等修饰语）
            out["appearance_spec"] = spec
            out["appearance_result"] = result
        elif key == "ph":
            out["ph_label"] = label or out["ph_label"]
            out["ph_spec"] = spec
            out["ph_result"] = result
        elif key == "solid":
            out["solid_label"] = label or out["solid_label"]
            out["solid_spec"] = spec
            out["solid_result"] = result
        elif key == "odour":
            out["odour_label"] = label or out["odour_label"]
            out["odour_spec"] = spec
            out["odour_result"] = result
    # 前端编辑后的扁平字段覆盖 tests[] 提取值
    for k in out:
        flat = batch.get(k)
        if flat not in (None, ""):
            out[k] = str(flat)
    return out


def _build_mapping(payload: Dict[str, Any]) -> Dict[str, str]:
    items = payload.get("items") or []
    first = items[0] if items else {}
    po_no = payload.get("po_no") or ""
    po_line = f"PO#{po_no}" if payload.get("show_po") and po_no else ""
    packages_display = payload.get("packages_display")
    if packages_display is None:
        packages_display = (
            payload.get("pallets") if payload.get("pallets") is not None else payload.get("packages", "")
        )
    raw: Dict[str, Any] = {
        "COMPANY_NAME_EN": payload.get("company_name_en", ""),
        "COMPANY_ADDR_EN": payload.get("company_addr_en", ""),
        "COMPANY_TEL_EN": payload.get("company_tel_en", ""),
        "CONSIGNEE_NAME": payload.get("consignee_name", ""),
        "CONSIGNEE_ADDR": payload.get("consignee_addr", ""),
        "CONSIGNEE_TAX": payload.get("consignee_tax", ""),
        "NOTIFY": payload.get("notify", ""),
        "NOTIFY_ADDR": payload.get("notify_addr", ""),
        "DEST_AGENT_NAME": payload.get("dest_agent_name", ""),
        "DEST_AGENT_ADDR": payload.get("dest_agent_addr", ""),
        "DEST_AGENT_TAX": payload.get("dest_agent_tax", ""),
        "DEST_AGENT_TEL": payload.get("dest_agent_tel", ""),
        "INVOICE_NO": payload.get("invoice_no", ""),
        "PACKING_NO": payload.get("packing_no", ""),
        "PI_NO": payload.get("pi_no", ""),
        "INVOICE_DATE": payload.get("invoice_date", ""),
        "PRICE_TERM": payload.get("price_term", ""),
        "PAYMENT_TERMS": payload.get("payment_terms", ""),
        "ROUTE": payload.get("route", ""),
        "LOADING_PORT": payload.get("loading_port", ""),
        "DISCHARGE_PORT": payload.get("discharge_port", ""),
        "VESSEL": payload.get("vessel", ""),
        "BL_NO": payload.get("bl_no", ""),
        "CONTAINER_NO": payload.get("container_no", ""),
        "SEAL_NO": payload.get("seal_no", ""),
        "ITEM_COUNT": len(items),
        "ITEM_DESC": first.get("desc", ""),
        "ITEM_QTY": _fmt_num(first.get("qty", "")),
        "ITEM_PRICE": _fmt_num(first.get("price", "")),
        "ITEM_AMOUNT": _fmt_num(first.get("amount", "")),
        "HS_CODES": payload.get("hs_codes", ""),
        "PO_LINE": po_line,
        "TOTALS_LINE": payload.get("totals_line", ""),
        "TOTAL_QTY": _fmt_num(payload.get("total_qty", "")),
        "TOTAL_AMOUNT": _fmt_num(payload.get("total_amount", "")),
        "AMOUNT_WORDS": payload.get("amount_words", ""),
        "PACKAGES": _fmt_num(packages_display) if packages_display not in (None, "") else "",
        "VOLUME_CBM": _fmt_num(payload.get("volume_cbm", "")) if payload.get("volume_cbm", "") not in (None, "") else "",
        "NET_KG": _fmt_num(payload.get("net_kg", "")),
        "GROSS_KG": _fmt_num(payload.get("gross_kg", "")),
        "PRODUCT_NAME": payload.get("product_name", ""),
        "SHIPPED_QTY": payload.get("shipped_qty_text", ""),
        "SI_DESC_BLOCK": payload.get("si_desc_block", ""),
        "MARKS": payload.get("marks", "N/M"),
        "REMARK": payload.get("remark", ""),
        "PACKAGE_UNIT": payload.get("package_unit_label", ""),
        "FINAL_DEST": payload.get("final_dest", ""),
        "CONTAINER_QTY": payload.get("container_qty", ""),
        "PACKING_RANGE": payload.get("packing_range", ""),
        "BATCH_NO": payload.get("batch_no", ""),
        "PROD_DATE": payload.get("prod_date", ""),
        "EXP_DATE": payload.get("exp_date", ""),
        "COA_PI_NO": payload.get("coa_pi_no", ""),
        "PH_LABEL": payload.get("ph_label", ""),
        "PH_SPEC": payload.get("ph_spec", ""),
        "PH_RESULT": payload.get("ph_result", ""),
        "SOLID_LABEL": payload.get("solid_label", ""),
        "SOLID_SPEC": payload.get("solid_spec", ""),
        "SOLID_RESULT": payload.get("solid_result", ""),
        "APPEARANCE_SPEC": payload.get("appearance_spec", ""),
        "APPEARANCE_RESULT": payload.get("appearance_result", ""),
        "ODOUR_LABEL": payload.get("odour_label", ""),
        "ODOUR_SPEC": payload.get("odour_spec", ""),
        "ODOUR_RESULT": payload.get("odour_result", ""),
        "BANK_LINE1": payload.get("bank_line1", ""),
        "BANK_LINE2": payload.get("bank_line2", ""),
        "BANK_LINE3": payload.get("bank_line3", ""),
        "BANK_LINE4": payload.get("bank_line4", ""),
        "BANK_LINE5": payload.get("bank_line5", ""),
        "BANK_LINE6": payload.get("bank_line6", ""),
    }
    # Multi-item rows: ITEM_*_1..N (legacy ITEM_* stay as aliases of item 1)
    for idx, it in enumerate(items, 1):
        raw[f"ITEM_NO_{idx}"] = it.get("index", idx)
        raw[f"ITEM_DESC_{idx}"] = it.get("desc", "")
        raw[f"ITEM_QTY_{idx}"] = _fmt_num(it.get("qty"))
        raw[f"ITEM_PRICE_{idx}"] = _fmt_num(it.get("price"))
        raw[f"ITEM_AMOUNT_{idx}"] = _fmt_num(it.get("amount"))
        raw[f"ITEM_HS_{idx}"] = it.get("hs_code", "")
        raw[f"ITEM_PACKAGES_{idx}"] = _fmt_num(it.get("packages_display", it.get("drums", ""))) if it.get("packages_display") or it.get("drums") else ""
        raw[f"ITEM_CBM_{idx}"] = _fmt_num(it.get("cbm", "")) if it.get("cbm") else ""
        raw[f"ITEM_NET_{idx}"] = it.get("net_kg", "")
        raw[f"ITEM_GROSS_{idx}"] = it.get("gross_kg", "")
        raw[f"ITEM_PACKING_RANGE_{idx}"] = it.get("packing_range", "")
    return {k: _fmt(v) for k, v in raw.items()}


def _fill_workbook_on_sheet(ws, mapping: Dict[str, str]) -> None:
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            text = str(cell.value)
            if "{{" not in text:
                continue
            for key, val in mapping.items():
                text = text.replace("{{" + key + "}}", val)
            text = _PLACEHOLDER_RE.sub("", text)
            # 空值标签行（HS CODE: / QTY:KGS）整行清掉
            if text.strip() in ("HS CODE:", "HS CODE：", "QTY:KGS", "NET WEIGHT:KGS", "H.S.CODE:", "H.S. Code :"):
                cell.value = None
                continue
            cell.value = _coerce_number(text)


def _fill_workbook(wb, mapping: Dict[str, str]) -> None:
    for ws in wb.worksheets:
        _fill_workbook_on_sheet(ws, mapping)


def _unmerge_from_row(ws, from_row: int) -> list:
    """Unmerge ranges starting at from_row; return (min_col, min_row, max_col, max_row) list."""
    moves = []
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row >= from_row:
            moves.append((mr.min_col, mr.min_row, mr.max_col, mr.max_row))
            ws.unmerge_cells(str(mr))
    return moves


def _apply_merged_ranges(ws, moves: list, n: int) -> None:
    for min_col, min_row, max_col, max_row in moves:
        ws.merge_cells(
            start_row=min_row + n,
            start_column=min_col,
            end_row=max_row + n,
            end_column=max_col,
        )


def _find_row_containing(ws, token: str) -> Optional[int]:
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and token in str(cell.value):
                return cell.row
    return None


def _expand_ci_detail_rows(ws, n_items: int, payload: Dict[str, Any]) -> None:
    """CI 多产品：按成品版式「品名行 + QTY 行 + HS CODE 行」× N 组扩行。

    对齐 CI HT260721A01 / BLUETEC Invoice：每个产品下方跟 NET/QTY 与 H.S.CODE，
    不再只留一条汇总 QTY/HS 尾巴（船务反馈「有点混乱」）。
    """
    if n_items <= 1:
        return
    detail_row = _find_row_containing(ws, "{{ITEM_DESC}}")
    if detail_row is None:
        detail_row = _find_row_containing(ws, "{{ITEM_DESC_1}}") or 18

    block = 3  # 品名 + QTY + HS
    n_new = block * (n_items - 1)
    insert_at = detail_row + block
    merges = _unmerge_from_row(ws, insert_at)
    ws.insert_rows(insert_at, n_new)
    _apply_merged_ranges(ws, merges, n_new)

    src_height = ws.row_dimensions[detail_row].height
    items = payload.get("items") or []
    for i in range(n_items):
        base = detail_row + i * block
        for j in range(block):
            row = base + j
            if i > 0 or j > 0:
                src = detail_row + j
                if src_height and j == 0:
                    ws.row_dimensions[row].height = src_height
                for col in range(1, 7):
                    ws.cell(row, col)._style = copy.copy(ws.cell(src, col)._style)
        idx = i + 1
        it = items[i] if i < len(items) else {}
        ws.cell(base, 1).value = str(idx)
        ws.cell(base, 2).value = f"{{{{ITEM_DESC_{idx}}}}}"
        ws.cell(base, 3).value = f"{{{{ITEM_QTY_{idx}}}}}"
        ws.cell(base, 4).value = f"{{{{ITEM_PRICE_{idx}}}}}"
        ws.cell(base, 5).value = f"{{{{ITEM_AMOUNT_{idx}}}}}"
        ws.cell(base + 1, 1).value = None
        ws.cell(base + 1, 2).value = f"QTY:{{{{ITEM_QTY_{idx}}}}}KGS"
        ws.cell(base + 1, 3).value = None
        ws.cell(base + 1, 4).value = None
        ws.cell(base + 1, 5).value = None
        ws.cell(base + 2, 1).value = None
        # 无 HS 的行整行留空，不打「HS CODE:」空尾巴
        if it.get("hs_code"):
            ws.cell(base + 2, 2).value = f"HS CODE:{{{{ITEM_HS_{idx}}}}}"
        else:
            ws.cell(base + 2, 2).value = None
        ws.cell(base + 2, 3).value = None
        ws.cell(base + 2, 4).value = None
        ws.cell(base + 2, 5).value = None

    # TOTAL 行：写值汇总全部明细（覆盖全量 item 行，避免 SUM 范围过期）
    total_row = None
    for row in ws.iter_rows(min_col=1, max_col=6):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip() in ("TOTAL:", "TOTAL"):
                total_row = cell.row
                break
        if total_row:
            break
    if total_row is None:
        total_row = _find_row_containing(ws, "{{TOTAL_AMOUNT}}")
    if total_row is not None:
        _set_cell_value(ws, total_row, 3, payload.get("total_qty", 0))
        amount = payload.get("total_amount")
        _set_cell_value(ws, total_row, 5, amount if amount else None)


def _expand_pl_detail_rows(ws, n_items: int, payload: Dict[str, Any]) -> None:
    """PL 多产品：「装箱号 + 品名 + 件/体积/净/毛」+ QTY + HS × N 组，TOTAL 全量汇总。"""
    if n_items <= 1:
        return
    detail_row = _find_row_containing(ws, "{{ITEM_DESC}}")
    if detail_row is None:
        detail_row = _find_row_containing(ws, "{{ITEM_DESC_1}}") or 17

    block = 3
    n_new = block * (n_items - 1)
    insert_at = detail_row + block
    merges = _unmerge_from_row(ws, insert_at)
    ws.insert_rows(insert_at, n_new)
    _apply_merged_ranges(ws, merges, n_new)

    items = payload.get("items") or []
    src_height = ws.row_dimensions[detail_row].height
    for i in range(n_items):
        base = detail_row + i * block
        for j in range(block):
            row = base + j
            if i > 0 or j > 0:
                src = detail_row + j
                if src_height and j == 0:
                    ws.row_dimensions[row].height = src_height
                for col in range(1, 7):
                    ws.cell(row, col)._style = copy.copy(ws.cell(src, col)._style)
        idx = i + 1
        it = items[i] if i < len(items) else {}
        packing_range = it.get("packing_range") or str(idx)
        ws.cell(base, 1).value = packing_range
        ws.cell(base, 2).value = f"{{{{ITEM_DESC_{idx}}}}}"
        ws.cell(base, 3).value = f"{{{{ITEM_PACKAGES_{idx}}}}}"
        ws.cell(base, 4).value = f"{{{{ITEM_CBM_{idx}}}}}"
        ws.cell(base, 5).value = f"{{{{ITEM_NET_{idx}}}}}"
        ws.cell(base, 6).value = f"{{{{ITEM_GROSS_{idx}}}}}"
        ws.cell(base + 1, 1).value = None
        ws.cell(base + 1, 2).value = f"QTY:{{{{ITEM_QTY_{idx}}}}}KGS"
        for col in (3, 4, 5, 6):
            ws.cell(base + 1, col).value = None
        ws.cell(base + 2, 1).value = None
        if it.get("hs_code"):
            ws.cell(base + 2, 2).value = f"HS CODE:{{{{ITEM_HS_{idx}}}}}"
        else:
            ws.cell(base + 2, 2).value = None
        for col in (3, 4, 5, 6):
            ws.cell(base + 2, col).value = None

    total_row = None
    for row in ws.iter_rows(min_col=1, max_col=6):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip() in ("TOTAL:", "TOTAL"):
                total_row = cell.row
                break
        if total_row:
            break
    if total_row is None:
        rows = [
            r
            for r in range(1, (ws.max_row or 1) + 1)
            if any(
                isinstance(ws.cell(r, c).value, str) and "{{GROSS_KG}}" in str(ws.cell(r, c).value)
                for c in range(1, 7)
            )
        ]
        if rows:
            total_row = rows[-1]
    if total_row is not None:
        _set_cell_value(
            ws,
            total_row,
            3,
            payload.get("packages_display", sum(it.get("packages_display") or 0 for it in items)),
        )
        _set_cell_value(
            ws,
            total_row,
            4,
            payload.get("volume_cbm", sum(it.get("cbm") or 0 for it in items)),
        )
        _set_cell_value(
            ws,
            total_row,
            5,
            payload.get("net_kg", sum(it.get("net_kg") or 0 for it in items)),
        )
        _set_cell_value(
            ws,
            total_row,
            6,
            payload.get("gross_kg", sum(it.get("gross_kg") or 0 for it in items)),
        )


def _set_cell_value(ws, row: int, col: int, value: Any) -> None:
    """Write to the anchor cell if (row, col) falls inside a merge."""
    cell = ws.cell(row, col)
    if cell.__class__.__name__ == "MergedCell":
        for mr in ws.merged_cells.ranges:
            if mr.min_row <= row <= mr.max_row and mr.min_col <= col <= mr.max_col:
                ws.cell(mr.min_row, mr.min_col).value = value
                return
    cell.value = value


def _append_extra_notes(wb, extra_notes) -> None:
    if not extra_notes:
        return
    notes = [str(n).strip() for n in extra_notes if n and str(n).strip()]
    if not notes:
        return
    ws = wb.worksheets[0]
    start = (ws.max_row or 1) + 2
    for i, note in enumerate(notes):
        c = ws.cell(start + i, 1, note)
        c.alignment = Alignment(wrap_text=True, vertical="center")


class ClearanceDocService:
    def load_template(
        self, doc_type: str, customer_code: Optional[str] = None
    ) -> openpyxl.Workbook:
        if doc_type not in _DOC_TYPES:
            raise ValueError(f"Unknown clearance doc_type: {doc_type}")
        blob, _src = template_service.load_template_bytes(customer_code, doc_type)
        return openpyxl.load_workbook(BytesIO(blob))

    def generate(
        self,
        doc_type: str,
        record: LedgerRecordResponse,
        company_code: Optional[str] = None,
        overrides: Optional[Union[Dict[str, Any], Any]] = None,
        customer_code: Optional[str] = None,
    ) -> Tuple[bytes, str, str]:
        """返回 (xlsx_bytes, doc_key, b64)。"""
        if doc_type not in _DOC_TYPES:
            raise ValueError(f"Unknown clearance doc_type: {doc_type}")
        if isinstance(overrides, dict):
            ov = overrides
        elif overrides is None:
            ov = {}
        else:
            ov = dict(overrides)

        cust = customer_code or ov.get("customer_code") or getattr(record, "customer_code", None)
        payload = build_clearance_payload(record, company_code, ov)
        wb = self.load_template(doc_type, customer_code=cust)

        # COA 多批次：一批一 sheet（录音：不同批次出货放同一 COA 不同 sheet）
        batches = ov.get("batches") or []
        if doc_type == "coa" and isinstance(batches, list) and len(batches) >= 1:
            self._fill_coa_batches(wb, payload, batches)
        else:
            items = payload.get("items") or []
            if doc_type == "ci" and len(items) > 1:
                _expand_ci_detail_rows(wb.worksheets[0], len(items), payload)
            elif doc_type == "pl" and len(items) > 1:
                _expand_pl_detail_rows(wb.worksheets[0], len(items), payload)
            mapping = _build_mapping(payload)
            _fill_workbook(wb, mapping)
            _append_extra_notes(wb, ov.get("extra_notes"))

        buf = BytesIO()
        wb.save(buf)
        content = buf.getvalue()
        doc_key = f"{doc_type}_{int(time.time())}"
        return content, doc_key, base64.b64encode(content).decode()

    def _fill_coa_batches(self, wb, base_payload: dict, batches: list) -> None:
        """用检测报告批次填充 COA；多批时一批一 sheet。PI 固定本票。"""
        template_ws = wb.worksheets[0]
        # 清空多余 sheet，按批次数复制干净模板（先复制再填，避免拷到已替换占位符的 sheet）
        while len(wb.worksheets) > 1:
            wb.remove(wb.worksheets[1])
        sheets = [template_ws]
        for _ in range(1, len(batches)):
            sheets.append(wb.copy_worksheet(template_ws))

        for i, batch in enumerate(batches):
            ws = sheets[i]
            ws.title = str(batch.get("batch_no") or f"COA-{i + 1}")[:31]
            fields = _extract_batch_test_fields(batch)

            p = dict(base_payload)
            p["batch_no"] = batch.get("batch_no") or p.get("batch_no") or ""
            if batch.get("quantity_text"):
                p["shipped_qty_text"] = batch["quantity_text"]
            # PROD_DATE：批次自带 → 批号推导；多批不继承单批基线日期
            prod = batch.get("prod_date") or ""
            if not prod and p["batch_no"]:
                from app.services.coa_date_service import parse_production_date_from_batch

                pd = parse_production_date_from_batch(p["batch_no"])
                if pd:
                    prod = pd.isoformat()
            p["prod_date"] = prod
            # 检测项：报告写什么就填什么（含外观「久置变深」）
            p["appearance_spec"] = fields.get("appearance_spec") or ""
            p["appearance_result"] = fields.get("appearance_result") or ""
            p["ph_label"] = fields.get("ph_label") or "PH VALUE"
            p["ph_spec"] = fields.get("ph_spec") or ""
            p["ph_result"] = fields.get("ph_result") or ""
            p["solid_label"] = fields.get("solid_label") or "SOLID CONTENT(%)"
            p["solid_spec"] = fields.get("solid_spec") or ""
            p["solid_result"] = fields.get("solid_result") or ""
            p["odour_label"] = fields.get("odour_label") or ""
            p["odour_spec"] = fields.get("odour_spec") or ""
            p["odour_result"] = fields.get("odour_result") or ""
            # 一票一 COA：PI 固定本票
            p["coa_pi_no"] = base_payload.get("pi_no") or p.get("coa_pi_no") or ""

            mapping = _build_mapping(p)
            _fill_workbook_on_sheet(ws, mapping)
