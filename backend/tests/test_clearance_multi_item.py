import pytest

from app.schemas.ledger import LedgerItemSchema, LedgerRecordResponse
from app.services.clearance_doc_service import _build_mapping
from app.services.clearance_fields import build_clearance_payload


def _item(
    code: str,
    desc: str,
    qty: float,
    price: float,
    amount: float,
    hs: str,
    net: float,
    gross: float,
    cbm: float,
    drums: int,
    pallets: int,
) -> LedgerItemSchema:
    return LedgerItemSchema(
        internal_code=code,
        product_en=desc,
        quantity_kg=qty,
        unit_price=price,
        total_amount=amount,
        hs_code=hs,
        drum_count=drums,
        pallet_count=pallets,
        net_weight_kg=net,
        gross_weight_kg=gross,
        volume_cbm=cbm,
    )


def make_single_record() -> LedgerRecordResponse:
    return LedgerRecordResponse(
        id=1,
        order_no="HT260720SZ",
        customer_code="WA318",
        consignee_name="CHEMZONE CO., LTD.",
        consignee_address="26/1D TRAN QUY CAP ST., HO CHI MINH CITY, VIETNAM",
        consignee_tel="TAX CODE: 0316699088",
        destination="HOCHIMINH(CAT LAI), VIETNAM",
        loading_port="NANSHA, CHINA",
        price_term="CIF HOCHIMINH",
        payment_terms="100% TT AFTER SHIPMENT 30 DAYS",
        pi_date="2026-07-20",
        currency="USD",
        items=[
            _item("FIX-HT016H", "FIXING AGENT HT-016H", 4000.0, 2.6, 10400.0, "340241", 4000.0, 4308.0, 7.92, 32, 8)
        ],
        status="saved",
    )


def make_three_item_record() -> LedgerRecordResponse:
    return LedgerRecordResponse(
        id=2,
        order_no="HT260801SZ",
        customer_code="WA318",
        consignee_name="CHEMZONE CO., LTD.",
        consignee_address="26/1D TRAN QUY CAP ST., HO CHI MINH CITY, VIETNAM",
        consignee_tel="TAX CODE: 0316699088",
        destination="HOCHIMINH(CAT LAI), VIETNAM",
        loading_port="NANSHA, CHINA",
        price_term="CIF HOCHIMINH",
        payment_terms="100% TT AFTER SHIPMENT 30 DAYS",
        pi_date="2026-08-01",
        currency="USD",
        items=[
            _item("A-001", "PRODUCT ALPHA", 1000.0, 1.5, 1500.0, "340241", 1000.0, 1080.0, 1.8, 8, 2),
            _item("B-002", "PRODUCT BETA", 2000.0, 2.0, 4000.0, "380894", 2000.0, 2160.0, 3.6, 16, 4),
            _item("C-003", "PRODUCT GAMMA", 500.0, 3.0, 1500.0, "290545", 500.0, 540.0, 0.9, 4, 1),
        ],
        status="saved",
    )


def test_three_item_mapping_has_indexed_keys():
    p = build_clearance_payload(record=make_three_item_record(), company_code="honghao", overrides={})
    m = _build_mapping(p)

    assert m["ITEM_COUNT"] == "3"
    for i, (desc, qty, price, amount, hs) in enumerate(
        [
            ("PRODUCT ALPHA", 1000.0, 1.5, 1500.0, "340241"),
            ("PRODUCT BETA", 2000.0, 2.0, 4000.0, "380894"),
            ("PRODUCT GAMMA", 500.0, 3.0, 1500.0, "290545"),
        ],
        start=1,
    ):
        assert m[f"ITEM_NO_{i}"] == str(i)
        assert m[f"ITEM_DESC_{i}"] == desc
        assert float(m[f"ITEM_QTY_{i}"]) == qty
        assert float(m[f"ITEM_PRICE_{i}"]) == price
        assert float(m[f"ITEM_AMOUNT_{i}"]) == amount
        assert m[f"ITEM_HS_{i}"] == hs

    # PL extras（默认件数=drums）
    assert m["ITEM_PACKAGES_1"] == "8"
    assert m["ITEM_PACKAGES_2"] == "16"
    assert m["ITEM_PACKAGES_3"] == "4"
    assert float(m["ITEM_CBM_1"]) == 1.8
    assert float(m["ITEM_NET_2"]) == 2000.0
    assert float(m["ITEM_GROSS_3"]) == 540.0

    # legacy aliases point at item 1
    assert m["ITEM_DESC"] == m["ITEM_DESC_1"]
    assert m["ITEM_QTY"] == m["ITEM_QTY_1"]
    assert m["ITEM_PRICE"] == m["ITEM_PRICE_1"]
    assert m["ITEM_AMOUNT"] == m["ITEM_AMOUNT_1"]

    # no item 4 keys
    assert "ITEM_DESC_4" not in m
    assert "ITEM_NO_4" not in m


def test_single_item_still_fills_legacy_aliases():
    p = build_clearance_payload(record=make_single_record(), company_code="honghao", overrides={})
    m = _build_mapping(p)

    assert m["ITEM_COUNT"] == "1"
    assert m["ITEM_DESC"] == "FIXING AGENT HT-016H"
    assert m["ITEM_QTY"] == "4000"
    assert m["ITEM_PRICE"] == "2.6"
    assert m["ITEM_AMOUNT"] == "10400"
    assert m["ITEM_NO_1"] == "1"
    assert m["ITEM_DESC_1"] == "FIXING AGENT HT-016H"
    assert m["ITEM_HS_1"] == "340241"
    assert m["HS_CODES"] == "340241"


def test_totals_sum_all_three_items():
    p = build_clearance_payload(record=make_three_item_record(), company_code="honghao", overrides={})
    assert p["total_qty"] == 3500.0
    assert p["total_amount"] == 7000.0
    assert p["net_kg"] == 3500.0
    assert p["gross_kg"] == 3780.0
    assert p["volume_cbm"] == pytest.approx(6.3)
    assert p["packages"] == 28
    assert p["pallets"] == 7
    assert p["totals_line"] == "TOTAL: 28 DRUMS PACKED ON 7 PALLETS"
    assert p["amount_words"] == "TOTAL USD SEVEN THOUSAND ONLY."


def test_package_unit_drums_vs_pallets():
    p_default = build_clearance_payload(record=make_three_item_record(), company_code="honghao", overrides={})
    m_default = _build_mapping(p_default)
    # default drums（对齐成品 PL 件数列）
    assert m_default["ITEM_PACKAGES_1"] == "8"
    assert m_default["ITEM_PACKAGES_2"] == "16"
    assert m_default["ITEM_PACKAGES_3"] == "4"

    p_drum = build_clearance_payload(
        record=make_three_item_record(),
        company_code="honghao",
        overrides={"package_unit": "drums"},
    )
    m_drum = _build_mapping(p_drum)
    assert m_drum["ITEM_PACKAGES_1"] == "8"
    assert m_drum["ITEM_PACKAGES_2"] == "16"
    assert m_drum["ITEM_PACKAGES_3"] == "4"
    assert m_drum["PACKAGES"] == "28"

    p_pal = build_clearance_payload(
        record=make_three_item_record(),
        company_code="honghao",
        overrides={"package_unit": "pallets"},
    )
    m_pal = _build_mapping(p_pal)
    assert m_pal["ITEM_PACKAGES_1"] == "2"
    assert m_pal["PACKAGES"] == "7"


def _all_text(content: bytes) -> str:
    import io

    import openpyxl

    parts = []
    for ws in openpyxl.load_workbook(io.BytesIO(content)).worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    parts.append(str(cell.value))
    return "\n".join(parts)


def test_generate_ci_three_items_row_expand():
    from app.services.clearance_doc_service import ClearanceDocService

    svc = ClearanceDocService()
    content, doc_key, _ = svc.generate("ci", make_three_item_record(), "honghao", {})
    assert doc_key.startswith("ci_")
    text = _all_text(content)
    assert "PRODUCT ALPHA" in text
    assert "PRODUCT BETA" in text
    assert "PRODUCT GAMMA" in text
    assert "{{" not in text


def test_generate_ci_three_items_totals():
    import io

    import openpyxl

    from app.services.clearance_doc_service import ClearanceDocService

    svc = ClearanceDocService()
    content, _, _ = svc.generate("ci", make_three_item_record(), "honghao", {})
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb.worksheets[0]

    # 3 detail rows + TOTAL values = sum of 3 items
    descs = []
    qtys = []
    amounts = []
    total_qty = None
    total_amount = None
    for row in ws.iter_rows(min_col=1, max_col=5):
        a, b, c, d, e = (cell.value for cell in row)
        if a == "TOTAL:" or b == "TOTAL:":
            total_qty = float(c if c is not None else b)
            total_amount = float(e)
        elif a in ("1", "2", "3") and b in ("PRODUCT ALPHA", "PRODUCT BETA", "PRODUCT GAMMA"):
            descs.append(b)
            qtys.append(float(c))
            amounts.append(float(e))

    assert descs == ["PRODUCT ALPHA", "PRODUCT BETA", "PRODUCT GAMMA"]
    assert qtys == [1000.0, 2000.0, 500.0]
    assert amounts == [1500.0, 4000.0, 1500.0]
    assert total_qty == sum(qtys) == 3500.0
    assert total_amount == sum(amounts) == 7000.0


def test_generate_ci_single_item_unchanged():
    import io

    import openpyxl

    from app.services.clearance_doc_service import ClearanceDocService

    svc = ClearanceDocService()
    content, _, _ = svc.generate("ci", make_single_record(), "honghao", {})
    text = _all_text(content)
    assert "FIXING AGENT HT-016H" in text
    assert "{{" not in text
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb.worksheets[0]
    # single product: one detail row only (row 18 in public template)
    assert any(ws.cell(r,2).value == "FIXING AGENT HT-016H" for r in range(1,40))
    assert float(next(ws.cell(r,3).value for r in range(1,40) if ws.cell(r,2).value == "FIXING AGENT HT-016H")) == 4000.0
    assert float(next(ws.cell(r,5).value for r in range(1,40) if ws.cell(r,2).value == "FIXING AGENT HT-016H")) == 10400.0


def test_generate_pl_three_items_row_expand():
    import io

    import openpyxl

    from app.services.clearance_doc_service import ClearanceDocService

    svc = ClearanceDocService()
    content, doc_key, _ = svc.generate("pl", make_three_item_record(), "honghao", {})
    assert doc_key.startswith("pl_")
    text = _all_text(content)
    assert "PRODUCT ALPHA" in text
    assert "PRODUCT BETA" in text
    assert "PRODUCT GAMMA" in text
    assert "{{" not in text

    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb.worksheets[0]
    descs = []
    pkgs = []
    cbms = []
    nets = []
    grosses = []
    total_pkg = total_cbm = total_net = total_gross = None
    for row in ws.iter_rows(min_col=1, max_col=6):
        a, b, c, d, e, f = (cell.value for cell in row)
        if a == "TOTAL:" or b == "TOTAL:":
            total_pkg = float(c)
            total_cbm = float(d)
            total_net = float(e)
            total_gross = float(f)
        elif b in ("PRODUCT ALPHA", "PRODUCT BETA", "PRODUCT GAMMA"):
            descs.append(b)
            pkgs.append(float(c))
            cbms.append(float(d))
            nets.append(float(e))
            grosses.append(float(f))

    assert descs == ["PRODUCT ALPHA", "PRODUCT BETA", "PRODUCT GAMMA"]
    # default package_unit=drums（对齐成品件数列）
    assert pkgs == [8.0, 16.0, 4.0]
    assert cbms == pytest.approx([1.8, 3.6, 0.9])
    assert nets == [1000.0, 2000.0, 500.0]
    assert grosses == [1080.0, 2160.0, 540.0]
    assert total_pkg == sum(pkgs) == 28.0
    assert total_cbm == pytest.approx(sum(cbms))
    assert total_net == sum(nets) == 3500.0
    assert total_gross == sum(grosses) == 3780.0


def test_generate_pl_single_item_unchanged():
    from app.services.clearance_doc_service import ClearanceDocService

    svc = ClearanceDocService()
    content, _, _ = svc.generate("pl", make_single_record(), "honghao", {})
    text = _all_text(content)
    assert "FIXING AGENT HT-016H" in text
    assert "{{" not in text

    import io

    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb.worksheets[0]
    assert any(ws.cell(r, 2).value == "FIXING AGENT HT-016H" for r in range(1, 40))
    item_row = next(r for r in range(1, 40) if ws.cell(r, 2).value == "FIXING AGENT HT-016H")
    assert float(ws.cell(item_row, 5).value) == 4000.0
    assert float(ws.cell(item_row, 6).value) == 4308.0


def test_generate_pl_package_unit_drums():
    import io

    import openpyxl

    from app.services.clearance_doc_service import ClearanceDocService

    svc = ClearanceDocService()
    content, _, _ = svc.generate(
        "pl", make_three_item_record(), "honghao", {"package_unit": "drums"}
    )
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb.worksheets[0]
    pkgs = []
    total_pkg = None
    for row in ws.iter_rows(min_col=1, max_col=3):
        a, b, c = (cell.value for cell in row)
        if a == "TOTAL:" or b == "TOTAL:":
            total_pkg = float(c)
        elif b in ("PRODUCT ALPHA", "PRODUCT BETA", "PRODUCT GAMMA"):
            pkgs.append(float(c))
    assert pkgs == [8.0, 16.0, 4.0]
    assert total_pkg == sum(pkgs) == 28.0
