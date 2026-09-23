"""客户定制清关模板 + extra_notes 回归测试。"""
import io

import openpyxl
import pytest

from app.database import SessionLocal, init_db
from app.models.customer_template import CustomerTemplate
from app.schemas.ledger import LedgerItemSchema, LedgerRecordResponse
from app.services import template_service
from app.services.clearance_doc_service import ClearanceDocService

CUSTOMER_MARK = "CUSTOMER_ONLY_MARK"
NOTE_MARK = "SPECIAL_NOTE_LINE_XYZ"


def make_record() -> LedgerRecordResponse:
    return LedgerRecordResponse(
        id=1,
        order_no="HT260720SZ",
        customer_code="WA318",
        consignee_name="CHEMZONE CO., LTD.",
        consignee_address="26/1D TRAN QUY CAP ST., HO CHI MINH CITY, VIETNAM",
        consignee_tel="TAX CODE: 0316699088",
        destination="HOCHIMINH(CAT LAI), VIETNAM",
        loading_port="NANSHA, CHINA",
        price_term="CIF HOCHIMINH",
        payment_terms="100% TT AFTER SHIPMENT 30 DAYS",
        pi_date="2026-07-20",
        currency="USD",
        items=[
            LedgerItemSchema(
                internal_code="FIX-HT016H",
                product_en="FIXING AGENT HT-016H",
                quantity_kg=4000.0,
                unit_price=2.6,
                total_amount=10400.0,
                hs_code="340241",
                drum_count=32,
                pallet_count=8,
                net_weight_kg=4000.0,
                gross_weight_kg=4308.0,
                volume_cbm=7.92,
            )
        ],
        status="saved",
    )


@pytest.fixture(autouse=True)
def _db():
    init_db()
    yield
    db = SessionLocal()
    try:
        db.query(CustomerTemplate).delete()
        db.commit()
    finally:
        db.close()


@pytest.fixture
def svc():
    return ClearanceDocService()


def _all_text(wb) -> str:
    parts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    parts.append(str(cell.value))
    return "\n".join(parts)


def _make_marked_template_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = CUSTOMER_MARK
    ws["A2"] = "{{INVOICE_NO}}"
    ws["A3"] = "{{PRODUCT_NAME}}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_generate_falls_back_to_public_when_no_customer_template(svc):
    """无客户模板时应回退公共模板（填充公共占位符，且不含客户标记）。"""
    content, doc_key, _ = svc.generate(
        "ci",
        make_record(),
        "honghao",
        {"invoice_date": "2026-08-01"},
        customer_code="NO_SUCH_CUST_PUBLIC_FB",
    )
    assert doc_key.startswith("ci_")
    text = _all_text(openpyxl.load_workbook(io.BytesIO(content)))
    # 公共 CI 模板填充结果
    assert "IN260720SZ" in text
    assert "HT260720SZ" in text
    assert "FIXING AGENT HT-016H" in text
    assert CUSTOMER_MARK not in text
    assert "{{" not in text


def test_save_as_customer_template_then_generate_uses_customer_blob(svc):
    """另存客户模板后，generate 应使用客户 xlsx（含 CUSTOMER_ONLY_MARK）。"""
    blob = _make_marked_template_bytes()
    saved = template_service.save_customer_template(
        customer_code="CUST_MARK_TEST",
        doc_type="ci",
        template_blob=blob,
        company_code="honghao",
    )
    assert saved["customer_code"] == "CUST_MARK_TEST"
    assert saved["doc_type"] == "ci"
    assert saved["version"] == 1

    content, _, _ = svc.generate(
        "ci", make_record(), "honghao", {}, customer_code="CUST_MARK_TEST"
    )
    text = _all_text(openpyxl.load_workbook(io.BytesIO(content)))
    assert CUSTOMER_MARK in text
    # 客户模板中的占位符仍被填充
    assert "IN260720SZ" in text
    assert "FIXING AGENT HT-016H" in text
    assert "{{" not in text


def test_extra_notes_appended(svc):
    """overrides.extra_notes 每行一句应追加到输出工作簿。"""
    content, _, _ = svc.generate(
        "ci",
        make_record(),
        "honghao",
        {"extra_notes": [NOTE_MARK, "ANOTHER_NOTE_LINE"]},
        customer_code="NO_SUCH_CUST_NOTES",
    )
    text = _all_text(openpyxl.load_workbook(io.BytesIO(content)))
    assert NOTE_MARK in text
    assert "ANOTHER_NOTE_LINE" in text


def test_load_template_bytes_source_labels():
    """load_template_bytes 无客户模板时返回 source=public，有则 customer。"""
    _blob, src = template_service.load_template_bytes("NO_SUCH_CUST_SRC", "ci")
    assert src == "public"

    marker_blob = _make_marked_template_bytes()
    template_service.save_customer_template(
        customer_code="CUST_SRC_TEST",
        doc_type="pl",
        template_blob=marker_blob,
    )
    blob3, src3 = template_service.load_template_bytes("CUST_SRC_TEST", "pl")
    assert src3 == "customer"
    assert blob3 == marker_blob
