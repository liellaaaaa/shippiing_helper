# -*- coding: utf-8 -*-
"""品质部检测报告（.docx）字段提取 — 供 COA 自动填充。

样例结构（WA318 / WA254）：
  段落：客户 / 品名 + 编号 / 批号 + 数量
  表格：项目 | 技术指标 | 检查结果
        外观 | … | …
        50g/L pH值 或 1%pH值 | … | …
        含固量（%） | … | …
一个 docx 可含多张表（多批次一文件）。
"""
from __future__ import annotations

import re
from typing import Any, Optional

from docx import Document
from io import BytesIO


def _map_item_name(cn_name: str) -> tuple[str, str]:
    """中文检测项 → (key, 英文标签)。key: appearance|odour|solid|ph|other"""
    name = (cn_name or "").strip()
    low = name.lower().replace(" ", "")
    if "外观" in name:
        return "appearance", "APPEARANCE"
    if "气味" in name or "odour" in low or "odor" in low:
        return "odour", "ODOUR"
    if "含固" in name or "solid" in low:
        return "solid", "SOLID CONTENT(%)"
    if "ph" in low or "ph值" in name.replace(" ", ""):
        # 条件写入占位：1% → PH (1%)；50g/L 或 20% → PH VALUE (20%)（对齐成品样例）
        if re.search(r"1\s*%", name):
            return "ph", "PH (1%)"
        if re.search(r"50\s*g\s*/\s*L|20\s*%", name, re.I):
            return "ph", "PH VALUE (20%)"
        m = re.search(r"(\d+\s*%|\d+\s*g/L)", name, re.I)
        cond = (m.group(1) if m else "").replace(" ", "")
        return "ph", f"PH VALUE ({cond})" if cond else "PH VALUE"
    return "other", name


def parse_inspection_report(content: bytes) -> list[dict[str, Any]]:
    """解析检测报告 bytes → 每批次一个 dict。"""
    doc = Document(BytesIO(content))
    header: dict[str, Any] = {
        "customer": None,
        "product_name_cn": None,
        "product_code": None,
        "batch_no": None,
        "quantity_text": None,
    }
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if not t:
            continue
        if "客户" in t and "：" in t:
            header["customer"] = t.split("：", 1)[-1].strip() or header["customer"]
        if "品名" in t:
            m = re.search(r"品名[:：]\s*(.+?)(?:\s{2,}|编号|$)", t)
            if m:
                header["product_name_cn"] = m.group(1).strip()
        if "编号" in t:
            m = re.search(r"编号[:：]\s*(\S+)", t)
            if m:
                header["product_code"] = m.group(1)
        if "批号" in t:
            m = re.search(r"批号[:：]\s*(\S+)", t)
            if m:
                header["batch_no"] = m.group(1)
        if "数量" in t:
            m = re.search(r"数量[:：]\s*(\S+)", t)
            if m:
                header["quantity_text"] = m.group(1)

    batches: list[dict[str, Any]] = []
    tables = doc.tables or []
    if not tables:
        batches.append({**header, "tests": []})
        return batches

    for idx, table in enumerate(tables):
        meta = dict(header)
        if len(tables) > 1:
            # 多表多批次：批号可能写在各段落，按出现顺序对应
            batch_nos = re.findall(r"批号[:：]\s*(\S+)", "\n".join(p.text for p in doc.paragraphs))
            qtys = re.findall(r"数量[:：]\s*(\S+)", "\n".join(p.text for p in doc.paragraphs))
            if idx < len(batch_nos):
                meta["batch_no"] = batch_nos[idx]
            if idx < len(qtys):
                meta["quantity_text"] = qtys[idx]

        tests: list[dict[str, str]] = []
        for row in table.rows[1:]:
            cells = [(c.text or "").strip().replace("\n", " ") for c in row.cells]
            if len(cells) < 3:
                continue
            name, spec, result = cells[0], cells[1], cells[2]
            if not name or name in ("项目", "结论", "备注"):
                continue
            key, en_label = _map_item_name(name)
            tests.append(
                {
                    "key": key,
                    "name_cn": name,
                    "label_en": en_label,
                    "spec": spec,
                    "result": result,
                }
            )
        batches.append({**meta, "tests": tests})
    return batches


def parse_inspection_docx_path(path: str) -> list[dict[str, Any]]:
    with open(path, "rb") as f:
        return parse_inspection_report(f.read())
